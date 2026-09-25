# -*- coding: utf-8 -*-
"""不调用外部平台或 AI 的后端稳定性测试。"""
import asyncio
import io
import time
import unittest
from unittest import mock

from fastapi import HTTPException

import app
import resolver
import quota_identity


class StabilityTests(unittest.TestCase):
    def test_grid_quota_links_device_and_network_identity(self):
        state = {"date": time.strftime("%Y-%m-%d"), "daily_cny": 0.0,
                 "clients": {}, "credits": {}, "orders": {}}
        quota_identity.reserve(state, ["dev:a", "net:one"], 0.8, 1.0, 10.0, 0.2)
        quote = quota_identity.quote(state, ["dev:b", "net:one"], 0.22, 1.0, 10.0, 0.2)
        self.assertFalse(quote["allowed"])
        self.assertEqual(quote["used_cny"], 0.8)
        quote = quota_identity.quote(state, ["dev:a", "net:two"], 0.22, 1.0, 10.0, 0.2)
        self.assertFalse(quote["allowed"])

    def test_grid_free_experience_stops_before_exceeding_one_yuan(self):
        state = {"date": app.time.strftime("%Y-%m-%d"), "daily_cny": 0.0, "clients": {}}
        with mock.patch.object(app, "_grid_usage", state), \
             mock.patch.object(app, "_save_grid_usage_locked"):
            for _ in range(4):
                app._reserve_grid_cost("visitor", 0.24)
            self.assertAlmostEqual(state["clients"]["visitor"], 0.96)
            with self.assertRaises(HTTPException) as caught:
                app._reserve_grid_cost("visitor", 0.24)
            self.assertEqual(caught.exception.status_code, 402)
            self.assertEqual(caught.exception.detail["code"], "grid_payment_required")
            self.assertEqual(caught.exception.detail["payable_cny"], 0.34)

    def test_failed_grid_generation_releases_reserved_cost(self):
        state = {"date": app.time.strftime("%Y-%m-%d"), "daily_cny": 0.0, "clients": {}}
        with mock.patch.object(app, "_grid_usage", state), \
             mock.patch.object(app, "_save_grid_usage_locked"):
            app._reserve_grid_cost("visitor", 0.22)
            app._release_grid_cost("visitor", 0.22)
            self.assertEqual(state["clients"]["visitor"], 0.0)
            self.assertEqual(state["daily_cny"], 0.0)

    def test_platform_vtt_subtitle_is_parsed_with_timestamps(self):
        parsed = resolver._parse_platform_subtitle(
            b"WEBVTT\n\n00:00:01.000 --> 00:00:03.000\nHello world\n\n"
            b"00:00:03.000 --> 00:00:05.500\nNext line\n", "vtt")
        self.assertEqual(parsed["text"], "Hello world Next line")
        self.assertEqual(parsed["segments"][0]["start_ms"], 1000)
        self.assertEqual(parsed["segments"][1]["end_ms"], 5500)

    def test_platform_manual_subtitle_wins_over_auto_caption(self):
        info = {
            "subtitles": {"zh-CN": [{"ext": "vtt", "data":
                "WEBVTT\n\n00:00:00.000 --> 00:00:01.000\n人工字幕\n"}]},
            "automatic_captions": {"zh": [{"ext": "vtt", "data":
                "WEBVTT\n\n00:00:00.000 --> 00:00:01.000\n自动字幕\n"}]},
        }
        subtitle = resolver._subtitle_from_info(info, {})
        self.assertEqual(subtitle["text"], "人工字幕")
        self.assertEqual(subtitle["source"], "平台字幕")

    def test_srt_parser_preserves_time_and_explicit_speaker(self):
        rows = app._parse_srt("1\n00:00:01,000 --> 00:00:03,200\n小知：别再手抄笔记了\n\n2\n00:00:04,000 --> 00:00:05,000\n没有人物名\n".encode())
        self.assertEqual(rows[0]["speaker"], "小知")
        self.assertEqual(rows[0]["start_ms"], 1000)
        self.assertEqual(rows[1]["speaker"], "待校准")

    def test_creative_quality_rejects_empty_storyboard(self):
        with self.assertRaises(app.HTTPException):
            app._validate_creative_phase("storyboard", {"storyboard": []})

    def test_creative_quality_rejects_incomplete_nine_grid(self):
        with self.assertRaises(app.HTTPException):
            app._validate_creative_phase("storyboard", {
                "storyboard": [{"group": 1, "time": "0-15s",
                                "panels": [{"visual": "只有一格"}]}]
            }, 1)

    def test_creative_quality_accepts_complete_nine_grid(self):
        result = app._validate_creative_phase("storyboard", {
            "storyboard": [{"group": 1, "time": "0-15s", "panels": [
                {"visual": f"连续画面{i}"} for i in range(1, 10)
            ]}]
        }, 1)
        self.assertEqual([x["panel"] for x in result["storyboard"][0]["panels"]],
                         list(range(1, 10)))

    def test_creative_prompts_must_cover_every_storyboard_group(self):
        with self.assertRaises(app.HTTPException):
            app._validate_creative_phase("prompts", {"video_prompts": [
                {"source_groups": [1], "prompt": "只覆盖第一组"}
            ]}, [1, 2])

    def test_creative_understanding_reuses_existing_analysis(self):
        result = app._understanding_from_analysis(
            {"chapters": [{"time": "00:00", "label": "开场钩子", "summary": "提出问题"}],
             "overall_summary": "先提痛点，再给方案"},
            [{"start_ms": 0, "end_ms": 1000, "speaker": "小知", "text": "开始吧"},
             {"start_ms": 1000, "end_ms": 2000, "speaker": "待校准", "text": "好的"}],
        )
        self.assertEqual(result["characters"][0]["name"], "小知")
        self.assertEqual(result["story_units"][0]["summary"], "提出问题")
        self.assertTrue(result["story_units"][0]["emotion_changes"])
        self.assertTrue(result["story_units"][0]["transition"])
        self.assertEqual(result["speaker_calibration"][1]["speaker"], "待确认")

    def test_creative_understanding_repairs_missing_emotion_transition_and_shot_size(self):
        incomplete = {"story_units": [{"unit": 1, "emotion_changes": [], "transition": ""}]}
        complete = {"story_units": [{"unit": 1, "emotion_changes": [{"from": "平静", "to": "紧张"}],
                                     "transition": "冲突升级"}],
                    "shot_analysis": [{"shot": 1, "shot_size": "近景", "composition": "居中",
                                       "action": "抬头", "visual_value": "强化反应"}]}
        with mock.patch("app._creative_asset_prompt", side_effect=[incomplete, complete]) as generate:
            result = app._creative_understanding(
                [{"shot": 1, "start": 0, "end": 3, "image": "data:image/jpeg;base64,QQ=="}],
                [], {"api_key": "test"}, {})
        self.assertEqual(generate.call_count, 2)
        self.assertEqual(result["shot_analysis"][0]["shot_size"], "近景")

    def test_creative_json_mode_retries_malformed_output(self):
        malformed = mock.Mock(status_code=200)
        malformed.json.return_value = {
            "choices": [{"message": {"content": "not-json"}}]
        }
        success = mock.Mock(status_code=200)
        success.json.return_value = {
            "choices": [{"message": {"content": '{"storyboard":[{"panels":[{}],"time":"0-5s"}]}'}}]
        }
        with mock.patch("app.requests.post", side_effect=[malformed, success]) as post, \
             mock.patch("app.time.sleep"):
            result = app._creative_asset_prompt("请输出 JSON", [], {"api_key": "test", "model": "qwen3-vl-plus"})
        self.assertEqual(len(result["storyboard"]), 1)
        self.assertEqual(post.call_count, 2)
        sent = post.call_args_list[0].kwargs["json"]
        self.assertEqual(sent["response_format"], {"type": "json_object"})
        self.assertFalse(sent["enable_thinking"])

    def test_script_without_people_is_valid_for_product_only_video(self):
        result = app._validate_creative_phase("script", {
            "script": {"story_units": [{"unit": 1}], "shots": [{"shot": 1}]}
        })
        self.assertEqual(result["role_profiles"], [])
        self.assertEqual(result["product_profiles"], [])

    def test_explicit_speaker_evidence_is_not_overwritten_by_zero_confidence(self):
        segments = [{"start_ms": 1000, "end_ms": 2500, "text": "你好",
                     "speaker": "人物甲", "speaker_source": "asr"}]
        understanding = {"speaker_calibration": [{"start_ms": 1000, "end_ms": 2500,
                                                    "speaker": "待确认", "confidence": 0.0}]}
        result = app._apply_explicit_speaker_evidence(segments, understanding)
        row = result["speaker_calibration"][0]
        self.assertEqual(row["speaker"], "人物甲")
        self.assertGreaterEqual(row["confidence"], 0.9)
        self.assertEqual(row["evidence"], "语音模型声纹分离")

    def test_script_cannot_reference_unprovided_products_or_assets(self):
        result = app._sanitize_creative_script({
            "product_profiles": [{"asset_id": "made-up", "name": "虚构产品"}],
            "script": {
                "story_units": [{"unit": 1, "product_placement": "植入虚构产品"}],
                "shots": [{"shot": 1, "asset_ids": ["person-1", "made-up"]}],
            },
        }, [{"id": "person-1", "type": "person", "hidden": False}])
        self.assertEqual(result["product_profiles"], [])
        self.assertEqual(result["script"]["story_units"][0]["product_placement"], "")
        self.assertEqual(result["script"]["shots"][0]["asset_ids"], ["person-1"])

    def test_script_workbench_builds_prompt_before_appending_rules(self):
        model_result = {
            "role_profiles": [], "product_profiles": [],
            "script": {"story_units": [{"unit": 1}], "shots": [{"shot": 1}]},
            "post_copy": {"titles": [], "body": "", "tags": []},
        }
        with mock.patch("app.load_config", return_value={"api_key": "test"}), \
             mock.patch("app.call_qwen_text_json", return_value=model_result) as generate:
            result = asyncio.run(app.creative_workbench({
                "phase": "script", "assets": [], "deconstruction": {},
                "analysis": {}, "target_total_seconds": 60,
            }, None))
        self.assertTrue(result["script"]["shots"])
        self.assertIn("真实性硬规则", generate.call_args.args[0])

    def test_health_is_dependency_free(self):
        self.assertEqual(app.health(), {"ok": True, "service": "video-insight"})

    def test_readiness_reports_all_required_dependencies(self):
        checks = app.readiness_checks()
        self.assertEqual(set(checks), {"api_key", "ffmpeg", "static"})
        self.assertTrue(checks["ffmpeg"])
        self.assertTrue(checks["static"])

    def test_production_api_docs_are_not_exposed(self):
        if app.DEPLOY_MODE:
            self.assertIsNone(app.app.docs_url)
            self.assertIsNone(app.app.openapi_url)

    def test_analysis_queue_times_out_cleanly(self):
        async def scenario():
            old_timeout = app.ANALYSIS_QUEUE_TIMEOUT
            app.ANALYSIS_QUEUE_TIMEOUT = 1
            held_slots = [await app.acquire_analysis_slot("link")
                          for _ in range(app.MAX_LINK_CONCURRENT)]
            try:
                with self.assertRaises(HTTPException) as caught:
                    await app.acquire_analysis_slot("link")
                self.assertEqual(caught.exception.status_code, 503)
            finally:
                for slot in held_slots:
                    slot.release()
                app.ANALYSIS_QUEUE_TIMEOUT = old_timeout

        asyncio.run(scenario())

    def test_private_network_urls_are_rejected(self):
        with self.assertRaises(resolver.ResolveError):
            resolver.validate_public_url("http://127.0.0.1/admin")
        with mock.patch("resolver.socket.getaddrinfo", return_value=[
            (None, None, None, None, ("10.0.0.8", 443)),
        ]):
            with self.assertRaises(resolver.ResolveError):
                resolver.validate_public_url("https://example.com/video")

    def test_public_url_is_allowed(self):
        with mock.patch("resolver.socket.getaddrinfo", return_value=[
            (None, None, None, None, ("93.184.216.34", 443)),
        ]):
            self.assertEqual(
                resolver.validate_public_url("https://example.com/video.mp4"),
                "https://example.com/video.mp4",
            )

    def test_model_retries_transient_statuses_then_succeeds(self):
        overloaded = mock.Mock(status_code=503, headers={})
        limited = mock.Mock(status_code=429, headers={"Retry-After": "0.25"})
        success = mock.Mock(status_code=200, headers={})
        success.json.return_value = {
            "choices": [{"message": {"content": '{"title":"恢复成功"}'}}]
        }
        with mock.patch("app.requests.post", side_effect=[overloaded, limited, success]) as post, \
             mock.patch("app.time.sleep") as sleep:
            result = app.call_qwen([], {"api_key": "test-key"})
        self.assertEqual(result["title"], "恢复成功")
        self.assertEqual(post.call_count, 3)
        self.assertEqual(sleep.call_count, 2)
        sent = post.call_args_list[0].kwargs["json"]
        self.assertEqual(sent["response_format"], {"type": "json_object"})
        self.assertFalse(sent["enable_thinking"])
        self.assertEqual(sent["model"], app.FAST_VISION_MODEL)

    def test_deep_analysis_keeps_configured_quality_model(self):
        self.assertEqual(
            app._analysis_model({"model": "qwen3-vl-plus"}, "deep"),
            "qwen3-vl-plus",
        )

    def test_model_does_not_retry_permanent_client_error(self):
        rejected = mock.Mock(status_code=400, headers={})
        with mock.patch("app.requests.post", return_value=rejected) as post, \
             mock.patch("app.time.sleep") as sleep:
            with self.assertRaises(HTTPException):
                app.call_qwen([], {"api_key": "test-key"})
        self.assertEqual(post.call_count, 1)
        sleep.assert_not_called()

    def test_light_and_heavy_jobs_use_separate_pools(self):
        async def scenario():
            link_slot = await app.acquire_analysis_slot("link")
            try:
                frame_slot = await app.acquire_analysis_slot("frames")
                frame_slot.release()
            finally:
                link_slot.release()

        asyncio.run(scenario())

    def test_agent_jobs_use_a_separate_pool(self):
        async def scenario():
            link_slot = await app.acquire_analysis_slot("link")
            try:
                agent_slot = await app.acquire_analysis_slot("agent")
                agent_slot.release()
            finally:
                link_slot.release()

        asyncio.run(scenario())

    def test_transcript_text_supports_whole_text_and_sentences(self):
        payload = {
            "transcripts": [
                {"text": "第一段完整转写。"},
                {"sentences": [{"text": "第二段。"}, {"text": "第三段。"}]},
            ]
        }
        self.assertEqual(
            app._transcript_text(payload),
            "第一段完整转写。\n第二段。第三段。",
        )

    def test_remote_asr_submits_polls_and_reads_transcript(self):
        submitted = mock.Mock(status_code=200)
        submitted.json.return_value = {"output": {"task_id": "task-1"}}
        completed = mock.Mock(status_code=200)
        completed.json.return_value = {
            "output": {
                "task_status": "SUCCEEDED",
                "result": {"transcription_url": "https://example.com/result.json"},
            }
        }
        transcript = mock.Mock(status_code=200)
        transcript.json.return_value = {
            "transcripts": [{"sentences": [{"text": "这是语音内容。"}]}]
        }
        transcript.raise_for_status.return_value = None
        with mock.patch("app.requests.post", return_value=submitted) as post, \
             mock.patch("app.requests.get", side_effect=[completed, transcript]) as get:
            text = app.transcribe_remote_audio(
                "https://example.com/video.mp4", {"api_key": "test-key"}, timeout=1)
        self.assertEqual(text, "这是语音内容。")
        self.assertEqual(post.call_args.kwargs["json"]["input"]["file_urls"],
                         ["https://example.com/video.mp4"])
        self.assertEqual(get.call_count, 2)

    def test_agent_plan_rejects_unknown_and_duplicate_tools(self):
        analysis = {"teaching": {"is_teaching": True}}
        plan = app._normalize_agent_plan({
            "intent": "整理课程",
            "route": "课程",
            "steps": [
                {"tool": "delete_everything", "reason": "越权工具"},
                {"tool": "course_pack", "reason": "生成课件"},
                {"tool": "course_pack", "reason": "重复调用"},
                {"tool": "creative_pack", "reason": "生成二创"},
            ],
        }, analysis, ["auto"])
        self.assertEqual([x["tool"] for x in plan["steps"]],
                         ["course_pack", "creative_pack"])

    def test_agent_plan_has_safe_fallback(self):
        analysis = {"teaching": {"is_teaching": False}}
        plan = app._normalize_agent_plan({}, analysis, ["auto"])
        self.assertEqual(plan["steps"], [{
            "tool": "creative_pack", "reason": "生成多平台二创素材",
        }])

    def test_agent_plan_respects_explicit_course_goal(self):
        analysis = {"teaching": {"is_teaching": True}}
        plan = app._normalize_agent_plan({
            "steps": [{"tool": "creative_pack", "reason": "模型偏好二创"}],
        }, analysis, ["course"])
        self.assertEqual(plan["steps"], [{
            "tool": "course_pack", "reason": "用户明确选择了课程整理路线",
        }])

    def test_agent_plan_all_goal_runs_both_bounded_tools(self):
        plan = app._normalize_agent_plan({}, {"teaching": {"is_teaching": True}}, ["all"])
        self.assertEqual([x["tool"] for x in plan["steps"]],
                         ["creative_pack", "course_pack"])

    def test_agent_quality_requires_complete_outputs(self):
        plan = {"steps": [{"tool": "creative_pack"}, {"tool": "course_pack"}]}
        incomplete = app._agent_quality({"creative": {"scripts": {"60s": "x"}}}, plan)
        self.assertFalse(incomplete["passed"])
        complete = app._agent_quality({
            "creative": {"scripts": {"60s": "x"}, "storyboard": [{}]},
            "course": {"outline": [{}], "slides": [{}]},
        }, plan)
        self.assertTrue(complete["passed"])

    def test_workbench_audit_blocks_incomplete_project(self):
        result = app._workbench_audit({
            "deconstruction": {"shots": [{"shot": 1}]},
            "assets": [{"id": "person-1", "type": "person"}],
            "script": {"script": {"shots": [{"visual": "人物走进房间", "dialogue": "你好"}]}},
            "storyboard": [{"group": 1, "panels": [{"visual": "只有一格"}]}],
            "prompts": [],
        })
        self.assertFalse(result["ready"])
        self.assertLess(result["score"], 100)
        self.assertTrue(any(x["name"] == "九宫格" and not x["passed"]
                            for x in result["checks"]))

    def test_workbench_audit_does_not_pass_zero_of_zero_dialogue(self):
        result = app._workbench_audit({})
        dialogue = next(x for x in result["checks"] if x["name"] == "人物台词")
        self.assertFalse(dialogue["passed"])
        self.assertEqual(dialogue["detail"], "0/0 个镜头有台词")

    def test_missing_overall_summary_is_grounded_in_existing_facts(self):
        result = app._ensure_overall_summary({
            "key_info": ["第一条事实", "第二条事实"],
            "speech_summary": "语音摘要",
            "overall_summary": "",
        })
        self.assertIn("第一条事实", result["overall_summary"])
        self.assertIn("语音摘要", result["overall_summary"])

    def test_workbench_agent_patches_are_field_limited(self):
        workbench = {
            "script": {"script": {"shots": [{"visual": "旧画面"}]}},
            "storyboard": [{"continuity_in": "旧承接"}],
        }
        patches = app._sanitize_agent_patches({"patches": [
            {"target": "script_shot", "index": 0, "field": "visual", "value": "新画面"},
            {"target": "script_shot", "index": 0, "field": "asset_ids", "value": "越权"},
            {"target": "storyboard_group", "index": 9, "field": "continuity_in", "value": "越界"},
        ]}, workbench)
        self.assertEqual(patches, [{"target": "script_shot", "index": 0,
                                    "field": "visual", "value": "新画面", "reason": ""}])

    def test_scenario_context_rejects_unknown_values(self):
        self.assertEqual(app._scenario_context("creative")[0], "creative")
        self.assertEqual(app._scenario_context("unknown")[0], "course")

    def test_native_pptx_and_pdf_exports(self):
        import base64
        from PIL import Image
        image_buffer = io.BytesIO()
        Image.new("RGB", (640, 360), (91, 92, 226)).save(image_buffer, format="PNG")
        report = {
            "title": "高效学习课程",
            "key_info": ["主动回忆", "间隔重复"],
            "overall_summary": "课程总结",
            "course": {
                "learning_objectives": ["掌握三个方法"],
                "outline": [{"title": "第一章", "points": ["理解主动回忆"]}],
                "slides": [{
                    "title": "主动回忆",
                    "bullets": ["合上书本自测"],
                    "speaker_notes": "请让学员现场练习",
                }],
            },
            "_export_images": ["data:image/png;base64," + base64.b64encode(image_buffer.getvalue()).decode()],
        }
        docx = app._build_docx(report).getvalue()
        pptx = app._build_pptx(report).getvalue()
        pdf = app._build_pdf(report).getvalue()
        self.assertTrue(docx.startswith(b"PK"))
        self.assertTrue(pptx.startswith(b"PK"))
        self.assertTrue(pdf.startswith(b"%PDF"))
        from docx import Document
        from pptx import Presentation
        document = Document(io.BytesIO(docx))
        self.assertIn("总体总结", "\n".join(p.text for p in document.paragraphs))
        self.assertEqual(len(document.inline_shapes), 1)
        deck = Presentation(io.BytesIO(pptx))
        self.assertGreaterEqual(len(deck.slides), 4)
        self.assertTrue(any(getattr(slide.shapes, "title", None) and
                            "主动回忆" in slide.shapes.title.text for slide in deck.slides))


if __name__ == "__main__":
    unittest.main()
